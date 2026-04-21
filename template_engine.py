"""
Движок шаблонов Word.
─────────────────────
Как работает:
  1. Пользователь создаёт .docx файл в папке templates/
  2. В тексте документа расставляет метки вида {{поле}}
  3. При генерации все метки заменяются реальными данными студента

Доступные метки (все поля студента):
  {{фамилия}}            {{имя}}               {{отчество}}
  {{фамилия_лат}}        {{имя_лат}}
  {{дата_рождения}}      {{день}}  {{месяц}}  {{год}}
  {{гражданство}}
  {{тип_документа}}      {{серия_документа}}   {{номер_документа}}
  {{документ_с}}         {{документ_по}}
  {{виза_серия}}         {{виза_номер}}
  {{виза_с}}             {{виза_по}}
  {{телефон}}            {{email}}
  {{форма_обучения}}     {{форма_оплаты}}
  {{адрес}}
  {{дата_прибытия}}
  {{работодатель}}       {{дата_работы}}
  {{сегодня}}            {{университет}}       {{цасиг}}
"""

from __future__ import annotations

import os
import shutil
from datetime import datetime
from pathlib import Path
from typing import Optional

try:
    from docx import Document
    from docx.oxml.ns import qn
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

TEMPLATES_DIR = Path(__file__).parent / "templates"
EXPORTS_DIR   = Path(__file__).parent / "exports"

TEMPLATES_DIR.mkdir(exist_ok=True)
EXPORTS_DIR.mkdir(exist_ok=True)


# ── построить словарь замен для студента ─────────────────────────────────────

def build_context(student: dict, country_name: str = "") -> dict:
    """Вернуть словарь  метка → значение  для данного студента."""

    def s(key):
        v = student.get(key)
        return str(v).strip() if v else ""

    def fmt_date(day, month, year):
        try:
            return f"{int(day):02d}.{int(month):02d}.{year}"
        except Exception:
            return ""

    birth = fmt_date(student.get("birth_day"),
                     student.get("birth_month"),
                     student.get("birth_year"))

    fio_ru  = f"{s('last_name_ru')} {s('first_name_ru')} {s('middle_name_ru')}".strip()
    fio_lat = f"{s('last_name_lat')} {s('first_name_lat')} {s('middle_name_lat')}".strip()

    return {
        "{{фио}}":              fio_ru,
        "{{фамилия}}":          s("last_name_ru"),
        "{{имя}}":              s("first_name_ru"),
        "{{отчество}}":         s("middle_name_ru"),
        "{{фамилия_лат}}":      s("last_name_lat"),
        "{{имя_лат}}":          s("first_name_lat"),
        "{{фио_лат}}":          fio_lat,
        "{{отчество_лат}}": s("middle_name_lat"),
        "{{дата_рождения}}":    birth,
        "{{день}}":             str(student.get("birth_day") or ""),
        "{{месяц}}":            str(student.get("birth_month") or ""),
        "{{год}}":              str(student.get("birth_year") or ""),
        "{{гражданство}}":      country_name,
        "{{тип_документа}}":    s("id_doc_type"),
        "{{серия_документа}}":  s("id_doc_series"),
        "{{номер_документа}}":  s("id_doc_number"),
        "{{документ_с}}":       s("id_doc_valid_from"),
        "{{документ_по}}":      s("id_doc_valid_to"),
        "{{виза_серия}}":       s("visa_series"),
        "{{виза_номер}}":       s("visa_number"),
        "{{виза_с}}":           s("visa_valid_from"),
        "{{виза_по}}":          s("visa_valid_to"),
        "{{телефон}}":          s("phone"),
        "{{email}}":            s("email"),
        "{{форма_обучения}}":   s("education_form"),
        "{{форма_оплаты}}":     s("payment_form"),
        "{{адрес}}":            s("registration_address"),
        "{{дата_прибытия}}":    s("arrival_date"),
        "{{работодатель}}":     s("employer_info"),
        "{{дата_работы}}":      s("employment_date"),
        # служебные
        "{{сегодня}}":          datetime.now().strftime("%d.%m.%Y"),
        "{{университет}}":      "ННГАСУ",
        "{{цасиг}}":            "ЦАСИГ",
        "{{руководитель}}":     "Белоус Е.А.",
    }


# ── замена меток в тексте ─────────────────────────────────────────────────────

def _replace_in_run(run, context: dict):
    for marker, value in context.items():
        if marker in run.text:
            run.text = run.text.replace(marker, value)


def _replace_in_paragraph(para, context: dict):
    """
    Заменить метки в параграфе.
    Сначала пробуем простую замену в каждом run,
    затем — если метка разбита между runs — склеиваем текст параграфа.
    """
    # Быстрый путь: каждый run содержит метку целиком
    for run in para.runs:
        _replace_in_run(run, context)

    # Медленный путь: метка разбита между runs
    full_text = "".join(r.text for r in para.runs)
    needs_merge = any(m in full_text for m in context)
    if needs_merge:
        for marker, value in context.items():
            if marker in full_text:
                full_text = full_text.replace(marker, value)
        # записать в первый run, остальные очистить
        if para.runs:
            para.runs[0].text = full_text
            for r in para.runs[1:]:
                r.text = ""


def _replace_in_table(table, context: dict):
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                _replace_in_paragraph(para, context)
            for nested_table in cell.tables:
                _replace_in_table(nested_table, context)


def fill_template(template_path: str, student: dict,
                  country_name: str = "") -> Optional[str]:
    """
    Заполнить шаблон данными студента.
    Вернуть путь к готовому файлу или None при ошибке.
    """
    if not DOCX_AVAILABLE:
        return None

    context = build_context(student, country_name)

    doc = Document(template_path)

    # параграфы верхнего уровня
    for para in doc.paragraphs:
        _replace_in_paragraph(para, context)

    # таблицы
    for table in doc.tables:
        _replace_in_table(table, context)

    # колонтитулы
    for section in doc.sections:
        for para in section.header.paragraphs:
            _replace_in_paragraph(para, context)
        for para in section.footer.paragraphs:
            _replace_in_paragraph(para, context)

    # имя выходного файла
    tpl_name = Path(template_path).stem
    last_name = (student.get("last_name_ru") or "student").replace(" ", "_")
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = EXPORTS_DIR / f"{tpl_name}_{last_name}_{ts}.docx"

    doc.save(str(out_path))
    return str(out_path)


# ── управление шаблонами ──────────────────────────────────────────────────────

def get_all_templates() -> list[dict]:
    """Вернуть список шаблонов из папки templates/."""
    result = []
    for f in sorted(TEMPLATES_DIR.glob("*.docx")):
        result.append({
            "name": f.stem,
            "filename": f.name,
            "path": str(f),
            "size": f.stat().st_size,
            "modified": datetime.fromtimestamp(f.stat().st_mtime).strftime("%d.%m.%Y %H:%M"),
        })
    return result


def add_template(src_path: str) -> tuple[bool, str]:
    """Скопировать файл шаблона в папку templates/."""
    src = Path(src_path)
    if not src.exists():
        return False, "Файл не найден"
    if src.suffix.lower() != ".docx":
        return False, "Только .docx файлы"
    dst = TEMPLATES_DIR / src.name
    if dst.exists():
        return False, f"Шаблон «{src.name}» уже существует"
    shutil.copy2(str(src), str(dst))
    return True, f"Шаблон «{src.name}» добавлен"


def delete_template(filename: str) -> tuple[bool, str]:
    """Удалить шаблон из папки templates/."""
    path = TEMPLATES_DIR / filename
    if not path.exists():
        return False, "Файл не найден"
    path.unlink()
    return True, f"Шаблон «{filename}» удалён"


def open_templates_folder():
    """Открыть папку templates/ в проводнике."""
    import subprocess, platform
    if platform.system() == "Windows":
        os.startfile(str(TEMPLATES_DIR))
    elif platform.system() == "Darwin":
        subprocess.Popen(["open", str(TEMPLATES_DIR)])
    else:
        subprocess.Popen(["xdg-open", str(TEMPLATES_DIR)])


def create_example_template():
    """Создать пример шаблона если папка пустая."""
    if not DOCX_AVAILABLE:
        return
    if any(TEMPLATES_DIR.glob("*.docx")):
        return

    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(3)
        sec.right_margin = Cm(1.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ПРИМЕР ШАБЛОНА")
    r.bold = True
    r.font.size = Pt(14)

    doc.add_paragraph()
    doc.add_paragraph("Настоящим подтверждается, что студент {{фио}},")
    doc.add_paragraph("гражданин(ка) {{гражданство}}, дата рождения {{дата_рождения}},")
    doc.add_paragraph("документ: {{тип_документа}} №{{номер_документа}},")
    doc.add_paragraph("действителен до {{документ_по}},")
    doc.add_paragraph("является студентом {{университет}}.")
    doc.add_paragraph()
    doc.add_paragraph("Форма обучения: {{форма_обучения}}")
    doc.add_paragraph("Форма оплаты: {{форма_оплаты}}")
    doc.add_paragraph("Адрес: {{адрес}}")
    doc.add_paragraph("Телефон: {{телефон}}")
    doc.add_paragraph()
    doc.add_paragraph("Дата выдачи: {{сегодня}}")
    doc.add_paragraph()
    p2 = doc.add_paragraph()
    p2.add_run("Руководитель {{цасиг}}").bold = True
    p2.add_run("          {{руководитель}}")

    doc.save(str(TEMPLATES_DIR / "пример_шаблона.docx"))
